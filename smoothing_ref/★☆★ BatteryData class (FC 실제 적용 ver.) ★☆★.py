import numpy as np
import pandas as pd
import xlwings as xw
from openpyxl.utils import get_column_letter as get_char
import matplotlib.pyplot as plt
import time

from Lib_LKS_common import ActCell, from_excel2
from Lib_LKS_units import UNIT_check; import astropy.units as u
from Lib_LKS_BatteryData import BatteryData
import warnings
warnings.simplefilter("ignore", category=RuntimeWarning)

UNIT = UNIT_check() # {'t': u.sec, 'V': u.V, 'I': u.A, 'Q': u.Ah, 'T': u.deg_C3}

paths = [
# Charge
         'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_SDI_PA2/23T/23T_SDI_PA2_CHG_CC_0.2.dat',
         'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_SDI_PA2/23T/23T_SDI_PA2_CHG_CC_0.3.dat',
         'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_SDI_PA2/23T/23T_SDI_PA2_CHG_CC_0.5.dat',
         'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_SDI_PA2/23T/23T_SDI_PA2_CHG_CC_0.8.dat',
         'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_SDI_PA2/23T/23T_SDI_PA2_CHG_CC_1.0.dat',
         'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_SDI_PA2/23T/23T_SDI_PA2_CHG_CC_1.2.dat',
         'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_SDI_PA2/23T/23T_SDI_PA2_CHG_CC_1.4.dat',
         'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_SDI_PA2/23T/23T_SDI_PA2_CHG_CC_1.6.dat',
         'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_SDI_PA2/23T/23T_SDI_PA2_CHG_CC_1.8.dat',
         'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_SDI_PA2/23T/23T_SDI_PA2_CHG_CC_2.0.dat',
# Discharge
         'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_SDI_PA2/23T/23T_SDI_PA2_DCH_CC_0.2.dat',
         'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_SDI_PA2/23T/23T_SDI_PA2_DCH_CC_0.1.dat',
         'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_SDI_PA2/23T/23T_SDI_PA2_DCH_CC_0.5.dat',
         'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_SDI_PA2/23T/23T_SDI_PA2_DCH_CC_1.0.dat',
        ]
# peaks = [4,3,2,2,2,2,2,2,2,2,
#          4,3,2,2]

# paths = [
# # Charge
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_SDI_PA2/45T/45T_SDI_PA2_CHG_CC_0.2.dat',
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_SDI_PA2/45T/45T_SDI_PA2_CHG_CC_0.5.dat',
# # Discharge
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_SDI_PA2/45T/45T_SDI_PA2_DCH_CC_0.2.dat',
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_SDI_PA2/45T/45T_SDI_PA2_DCH_CC_0.5.dat',
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_SDI_PA2/45T/45T_SDI_PA2_DCH_CC_1.0.dat',]

# peaks = [4,2,
#          3,2,2]

# paths = [
# # Charge
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_ATL_PS/23T/23T_ATL_PS_CHG_CC_0.2.dat',
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_ATL_PS/23T/23T_ATL_PS_CHG_CC_0.3.dat',
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_ATL_PS/23T/23T_ATL_PS_CHG_CC_0.5.dat',
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_ATL_PS/23T/23T_ATL_PS_CHG_CC_0.8.dat',
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_ATL_PS/23T/23T_ATL_PS_CHG_CC_1.0.dat',
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_ATL_PS/23T/23T_ATL_PS_CHG_CC_1.2.dat',
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_ATL_PS/23T/23T_ATL_PS_CHG_CC_1.4.dat',
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_ATL_PS/23T/23T_ATL_PS_CHG_CC_1.6.dat',
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_ATL_PS/23T/23T_ATL_PS_CHG_CC_1.8.dat',
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_ATL_PS/23T/23T_ATL_PS_CHG_CC_2.0.dat',
# # Discharge
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_ATL_PS/23T/23T_ATL_PS_DCH_CC_0.1.dat',
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_ATL_PS/23T/23T_ATL_PS_DCH_CC_0.2.dat',
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_ATL_PS/23T/23T_ATL_PS_DCH_CC_0.5.dat',
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_ATL_PS/23T/23T_ATL_PS_DCH_CC_1.0.dat',]
# peaks = [3,3,2,2,2,2,2,2,2,2,
#          3,3,2,2]

# paths = [
# # Charge
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_A55/23T/23T_A55_CHG_CC_0.2.dat',
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_A55/23T/23T_A55_CHG_CC_0.3.dat',
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_A55/23T/23T_A55_CHG_CC_0.5.dat',
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_A55/23T/23T_A55_CHG_CC_0.8.dat',
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_A55/23T/23T_A55_CHG_CC_1.0.dat',
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_A55/23T/23T_A55_CHG_CC_1.2.dat',
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_A55/23T/23T_A55_CHG_CC_1.4.dat',
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_A55/23T/23T_A55_CHG_CC_1.6.dat',
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_A55/23T/23T_A55_CHG_CC_1.8.dat',
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_A55/23T/23T_A55_CHG_CC_2.0.dat',
# # Discharge
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_A55/23T/23T_A55_DCH_CC_0.1.dat',
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_A55/23T/23T_A55_DCH_CC_0.2.dat',
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_A55/23T/23T_A55_DCH_CC_0.5.dat',
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_A55/23T/23T_A55_DCH_CC_1.0.dat',]
# peaks = [3,3,2,2,2,2,2,2,2,2,
#          3,3,2,2]

# paths = [
# # Charge
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_ATL_PS/45T/45T_ATL_PS_CHG_CC_0.2.dat',
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_ATL_PS/45T/45T_ATL_PS_CHG_CC_0.5.dat',
# # Discharge
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_ATL_PS/45T/45T_ATL_PS_DCH_CC_0.2.dat',
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_ATL_PS/45T/45T_ATL_PS_DCH_CC_0.5.dat',
#          'D:/work/26/Lib_LKS/260107) Lib_LKS_matching & sample/Converted/FC_ATL_PS/45T/45T_ATL_PS_DCH_CC_1.0.dat',]
#
# peaks = [4,2,
#          3,2,2]

dQdVs = []
dVdQs = []
Crates = []
# for path,peak in zip(paths,peaks):
for path in paths:
    df = pd.read_csv(path)
    path_part2 = path.split('/')[-1]
    print(path_part2)
    Crate = float(path.split('_')[-1].split('.dat')[0].split('.dat')[0])
    Crates.append(Crate)
    if (('CHG' in path) &
        (len(np.where(np.round(df.iloc[2:, 2], 2) < np.round(np.mean(df.iloc[2:12, 2]), 2))[0]) > 0)):
        df = df.iloc[:np.where((np.round(df.iloc[2:, 2], 2) < np.round(np.mean(df.iloc[2:12, 2]), 2))&
                               (np.round(df.iloc[:,1],2) >= np.round(np.max(df.iloc[:,1]),2)))[0][0], :]
    BD = BatteryData(df)
    peaks = BD.run(denoise_strength=3.5, Crate=Crate, slope_window=1)
    BD.df.to_excel(path.split('.dat')[0]+'.xlsx',index=False)

    # import openpyxl
    # from openpyxl import Workbook
    # from openpyxl.chart import ScatterChart, Reference, Series
    #
    # bk = openpyxl.load_workbook(path.split('.dat')[0] + '.xlsx')
    # sh = bk.active
    #
    # max_row = BD.df.shape[0] + 1
    # # find_col_V_denoise = np.where(BD.df.columns == 'V_denoise')[0][0] + 1
    # # find_col_Q_denoise = np.where(BD.df.columns == 'Q')[0][0] + 1
    # # find_col_dVdQ   = np.where(BD.df.columns == 'dVdQ')[0][0] + 1
    # # find_col_dVdQ_g = np.where(BD.df.columns == 'dVdQ_g')[0][0] + 1
    # # find_col_dQdV   = np.where(BD.df.columns == 'dQdV')[0][0] + 1
    # # find_col_dQdV_g = np.where(BD.df.columns == 'dQdV_g')[0][0] + 1
    # # find_col_dVdQ_m = np.where(BD.df.columns == 'dVdQ_m3')[0][0] + 1
    # # find_col_dQdV_m = np.where(BD.df.columns == 'dQdV_m3')[0][0] + 1
    # find_col_V_denoise = np.where(BD.df.columns == 'V')[0][0] + 1
    # find_col_Q_denoise = np.where(BD.df.columns == 'Q')[0][0] + 1
    # find_col_dVdQ   = np.where(BD.df.columns == 'dV/dQ')[0][0] + 1
    # find_col_dVdQ_g = np.where(BD.df.columns == 'dV/dQ_g')[0][0] + 1
    # find_col_dQdV   = np.where(BD.df.columns == 'dQ/dV')[0][0] + 1
    # find_col_dQdV_g = np.where(BD.df.columns == 'dQ/dV_g')[0][0] + 1
    # find_col_dVdQ_m = np.where(BD.df.columns == 'dV/dQ_m3')[0][0] + 1
    # find_col_dQdV_m = np.where(BD.df.columns == 'dQ/dV_m3')[0][0] + 1
    #
    # x_val_chart1 = Reference(sh, min_col=find_col_Q_denoise, min_row=2, max_row=max_row)
    # x_val_chart3 = Reference(sh, min_col=find_col_V_denoise, min_row=2, max_row=max_row)
    #
    # y_val_chart1_series1 = Reference(sh, min_col=find_col_dVdQ, min_row=2, max_row=max_row)
    # y_val_chart1_series2 = Reference(sh, min_col=find_col_dVdQ_g, min_row=2, max_row=max_row)
    # y_val_chart1_series3 = Reference(sh, min_col=find_col_dVdQ_m, min_row=2, max_row=max_row)
    #
    # y_val_chart2_series1 = Reference(sh, min_col=find_col_dQdV, min_row=2, max_row=max_row)
    # y_val_chart2_series2 = Reference(sh, min_col=find_col_dQdV_g, min_row=2, max_row=max_row)
    # y_val_chart2_series3 = Reference(sh, min_col=find_col_dQdV_m, min_row=2, max_row=max_row)
    #
    # chart1 = ScatterChart()
    # # chart1.style = 1
    # chart1.height = 7.8
    # chart1.width = 12.84
    # chart1.auto_axis = True
    # chart1.y_axis.scaling.min = 0.3 * (BD.type - 1)
    # chart1.y_axis.scaling.max = 0.3 * (BD.type)
    # chart1.series.append(Series(y_val_chart1_series1, x_val_chart1, title='dVdQ'))
    # chart1.series.append(Series(y_val_chart1_series2, x_val_chart1, title='dVdQ_g'))
    # chart1.series.append(Series(y_val_chart1_series3, x_val_chart1, title='dVdQ_m'))
    #
    # chart2 = ScatterChart()
    # # chart2.style = 1
    # chart2.height = 7.8
    # chart2.width = 12.84
    # chart2.auto_axis = True
    # chart2.y_axis.scaling.min = 30 * (BD.type - 1)
    # chart2.y_axis.scaling.max = 30 * (BD.type)
    # chart2.series.append(Series(y_val_chart2_series1, x_val_chart1, title='dVdQ'))
    # chart2.series.append(Series(y_val_chart2_series2, x_val_chart1, title='dVdQ_g'))
    # chart2.series.append(Series(y_val_chart2_series3, x_val_chart1, title='dVdQ_m'))
    #
    # chart3 = ScatterChart()
    # # chart3.style = 1
    # chart3.height = 7.8
    # chart3.width = 12.84
    # chart3.x_axis.scaling.min = np.floor(np.min(BD.arrays.get('V_denoise').value))
    # chart3.x_axis.scaling.max = np.ceil(np.max(BD.arrays.get('V_denoise').value))
    # chart3.x_axis.majorUnit = 0.1
    # chart3.y_axis.scaling.min = 0.3 * (BD.type - 1)
    # chart3.y_axis.scaling.max = 0.3 * (BD.type)
    # chart3.series.append(Series(y_val_chart1_series1, x_val_chart3, title='dQdV'))
    # chart3.series.append(Series(y_val_chart1_series2, x_val_chart3, title='dQdV_g'))
    # chart3.series.append(Series(y_val_chart1_series3, x_val_chart3, title='dQdV_m'))
    #
    # chart4 = ScatterChart()
    # # chart4.style = 1
    # chart4.height = 7.8
    # chart4.width = 12.84
    # chart4.x_axis.scaling.min = np.floor(np.min(BD.arrays.get('V_denoise').value))
    # chart4.x_axis.scaling.max = np.ceil(np.max(BD.arrays.get('V_denoise').value))
    # chart4.x_axis.majorUnit = 0.1
    # chart4.y_axis.scaling.min = 30 * (BD.type - 1)
    # chart4.y_axis.scaling.max = 30 * (BD.type)
    # chart4.series.append(Series(y_val_chart2_series1, x_val_chart3, title='dQdV'))
    # chart4.series.append(Series(y_val_chart2_series2, x_val_chart3, title='dQdV_g'))
    # chart4.series.append(Series(y_val_chart2_series3, x_val_chart3, title='dQdV_m'))
    #
    # sh.add_chart(chart1, f'{get_char(BD.df.shape[1] + 2)}2')
    # sh.add_chart(chart2, f'{get_char(BD.df.shape[1] + 2)}15')
    # sh.add_chart(chart3, f'{get_char(BD.df.shape[1] + 9)}2')
    # sh.add_chart(chart4, f'{get_char(BD.df.shape[1] + 9)}15')
    # time.sleep(1)
    #
    # bk.save(path.split('.dat')[0] + '.xlsx')
    # bk.close()